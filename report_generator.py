
from docx import Document
from io import BytesIO


def create_report(file_a, file_b, comparison):
    doc = Document()

    doc.add_heading("AI Document Comparison Report", 0)
    doc.add_paragraph(f"Document A: {file_a}")
    doc.add_paragraph(f"Document B: {file_b}")
    doc.add_heading("Comparison Results", level=1)

    for section in comparison.split("\n"):
        if section.strip():
            doc.add_paragraph(section)

    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return output.getvalue()
